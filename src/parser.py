import csv
import json
import logging
import random
import sys
import time
from pathlib import Path
from typing import Any

import requests
from bs4 import BeautifulSoup

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = "https://xn--80aackbd0bcms3a1b4gta.xn--p1ai"
LIST_PAGE = f"{BASE}/vacancy"
API_URL = f"{BASE}/api/vacancy/list"

DATA_DIR = Path("data")
CSV_PATH = DATA_DIR / "vacancies.csv"
XLSX_PATH = DATA_DIR / "vacancies.xlsx"
PROGRESS_PATH = DATA_DIR / "progress.json"

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                  "AppleWebKit/537.36 (KHTML, like Gecko) "
                  "Chrome/124.0 Safari/537.36",
    "Accept-Language": "ru-RU,ru;q=0.9",
}

log = logging.getLogger("parser")

FIELDS = [
    "url", "title", "speciality", "organization", "area",
    "work_mode", "vacancy_part", "qty", "updated",
    "mentors", "housing", "social_land", "social_communal",
    "zemskii", "social_rent", "social_mortgage", "social_deposit",
]

PAGE_SIZE = 60  # размер страницы, который запоминает сервер после POST формы #w0


# ---------- сессия и CSRF ----------

def _extract_csrf(html: str) -> str:
    meta = BeautifulSoup(html, "lxml").find("meta", {"name": "csrf-token"})
    if meta is None or not meta.get("content"):
        raise RuntimeError("CSRF-токен не найден на странице")
    return str(meta["content"])


def make_session(page_size: int = PAGE_SIZE) -> tuple[requests.Session, str]:
    """GET /vacancy -> cookies + csrf. POST формы #w0 -> pageSize в сессии."""
    s = requests.Session()
    s.headers.update(HEADERS)

    # 1) GET — получаем cookie сессии и csrf
    r = s.get(LIST_PAGE, timeout=30)
    r.raise_for_status()
    csrf = _extract_csrf(r.text)

    # 2) POST формы w0 — сервер запоминает размер страницы в сессии
    r = s.post(
        LIST_PAGE,
        data={"_csrf": csrf, "pageSize": str(page_size)},
        headers={"Referer": LIST_PAGE},
        timeout=30,
    )
    r.raise_for_status()

    # 3) перечитываем csrf — после POST он мог обновиться
    r = s.get(LIST_PAGE, timeout=30)
    r.raise_for_status()
    csrf = _extract_csrf(r.text)

    return s, csrf


# ---------- один запрос к API ----------

def fetch_page_raw(session: requests.Session, csrf: str, page: int) -> dict[str, Any]:
    data = {
        "_csrf": csrf,
        "page": str(page),
        "pageSize": str(PAGE_SIZE),
        "PVacancySearch[territory_name]": "",
        "PVacancySearch[territory]": "",
        "PVacancySearch[vacancy_category_name]": "",
        "PVacancySearch[vacancy_category]": "",
        "PVacancySearch[mentors]": "0",
        "PVacancySearch[organization_name]": "",
        "PVacancySearch[organization]": "",
        "PVacancySearch[vacancy_part]": "",
        "PVacancySearch[payment_main]": "",
        "PVacancySearch[payment_stimulating]": "",
        "PVacancySearch[payment_other]": "0",
        "PVacancySearch[housing]": "0",
        "PVacancySearch[social_land]": "0",
        "PVacancySearch[social_communal]": "0",
        "PVacancySearch[zemskii]": "0",
        "PVacancySearch[social_rent]": "0",
        "PVacancySearch[social_mortgage]": "0",
        "PVacancySearch[social_deposit]": "0",
    }
    r = session.post(
        API_URL,
        data=data,
        headers={
            "Accept": "application/json",
            "X-Requested-With": "XMLHttpRequest",
            "Referer": LIST_PAGE,
        },
        timeout=30,
    )
    r.raise_for_status()
    payload = r.json()
    if not isinstance(payload, dict):
        raise RuntimeError(f"API вернул не объект JSON: {type(payload)}")
    return payload


def fetch_page(session_ref: dict[str, Any], page: int,
               max_attempts: int = 5, page_size: int = PAGE_SIZE) -> dict[str, Any]:
    """session_ref — изменяемый словарь {'session': ..., 'csrf': ...}."""
    for attempt in range(1, max_attempts + 1):
        try:
            return fetch_page_raw(session_ref["session"], session_ref["csrf"], page)

        except requests.HTTPError as e:
            status = e.response.status_code if e.response is not None else None
            log.warning("HTTP %s на странице %s (попытка %s/%s)",
                        status, page, attempt, max_attempts)
            if status in (400, 403, 419):
                log.info("обновляю сессию и CSRF")
                session_ref["session"], session_ref["csrf"] = make_session(page_size)

        except (requests.ConnectionError,
                requests.Timeout,
                requests.exceptions.ChunkedEncodingError) as e:
            log.warning("сеть/соединение на странице %s (попытка %s/%s): %s",
                        page, attempt, max_attempts, e)
            try:
                session_ref["session"], session_ref["csrf"] = make_session(page_size)
            except Exception as ee:
                log.warning("не удалось пересоздать сессию: %s", ee)

        backoff = min(60, (2 ** attempt) + random.uniform(0, 1.5))
        log.info("жду %.1f сек перед повтором", backoff)
        time.sleep(backoff)

    raise RuntimeError(f"страница {page}: исчерпаны попытки")


# ---------- нормализация ----------

def item_to_row(item: dict[str, Any]) -> dict[str, Any]:
    cat = item.get("vacancyCategory") or {}
    cat_params = cat.get("catParams") or {}
    mo = item.get("mo") or {}
    area = item.get("area") or {}
    route = item.get("route") or "vacancy"
    alias = item.get("alias") or ""

    return {
        "url": f"{BASE}/{route}/{alias}" if alias else "",
        "title": item.get("name_alt") or cat.get("name") or "",
        "speciality": cat_params.get("name") or "",
        "organization": mo.get("short_name") or "",
        "area": area.get("name") or "",
        "work_mode": item.get("work_mode") or "",
        "vacancy_part": item.get("vacancy_part") or "",
        "qty": item.get("qty") or "",
        "updated": item.get("updated") or "",
        "mentors": item.get("mentors") or 0,
        "housing": item.get("housing") or 0,
        "social_land": item.get("social_land") or 0,
        "social_communal": item.get("social_communal") or 0,
        "zemskii": item.get("zemskii") or 0,
        "social_rent": item.get("social_rent") or 0,
        "social_mortgage": item.get("social_mortgage") or 0,
        "social_deposit": item.get("social_deposit") or 0,
    }


# ---------- сохранение ----------

def save_csv(rows: list[dict[str, Any]]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    tmp = CSV_PATH.with_suffix(".csv.tmp")
    with tmp.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)
    tmp.replace(CSV_PATH)

XLSX_COLUMN_WIDTHS = {
    "url": 55, "title": 40, "speciality": 25, "organization": 40,
    "area": 25, "work_mode": 25, "vacancy_part": 10, "qty": 8,
    "updated": 14, "mentors": 10, "housing": 10, "social_land": 10,
    "social_communal": 12, "zemskii": 10, "social_rent": 10,
    "social_mortgage": 10, "social_deposit": 10,
}


def save_xlsx(rows: list[dict[str, Any]]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    if ws is None:
        # теоретически невозможно у свежей книги, но Pylance требует проверки
        ws = wb.create_sheet("Вакансии")
    ws.title = "Вакансии"

    # Шапка
    ws.append(FIELDS)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1A3A5C")
    for i, col in enumerate(FIELDS, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = \
            XLSX_COLUMN_WIDTHS.get(col, 15)

    # Данные
    for r in rows:
        ws.append([r.get(f, "") for f in FIELDS])

    # Закрепление шапки и автофильтр
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    tmp = XLSX_PATH.with_suffix(".xlsx.tmp")
    wb.save(tmp)
    try:
        tmp.replace(XLSX_PATH)
    except PermissionError:
        # файл открыт в Excel — сохраняем с другим именем
        fallback = XLSX_PATH.with_name(XLSX_PATH.stem + "_new.xlsx")
        wb.save(fallback)
        log.warning("vacancies.xlsx занят (открыт в Excel?). "
                    "Сохранено как %s", fallback.name)

def load_progress() -> dict[str, Any]:
    if PROGRESS_PATH.exists():
        try:
            return json.loads(PROGRESS_PATH.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {"next_page": 1}


def save_progress(next_page: int, total_pages: int | None) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    PROGRESS_PATH.write_text(
        json.dumps({"next_page": next_page, "total_pages": total_pages},
                   ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def load_existing_rows() -> list[dict[str, Any]]:
    if not CSV_PATH.exists():
        return []
    with CSV_PATH.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


# ---------- основной цикл ----------

def scrape_all(page_size: int = PAGE_SIZE) -> list[dict[str, Any]]:
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    rows = load_existing_rows()
    seen = {(r.get("url") or r.get("title") or "") for r in rows}
    progress = load_progress()
    page = int(progress.get("next_page", 1))

    if rows:
        log.info("найдено сохранённых %s строк, продолжаю с страницы %s",
                 len(rows), page)

    session_ref: dict[str, Any] = {"session": None, "csrf": None}
    session_ref["session"], session_ref["csrf"] = make_session(page_size)
    log.info("сессия готова, pageSize=%s зафиксирован", page_size)

    first_iteration = True

    while True:
        log.info("запрос страницы %s (уже собрано %s)", page, len(rows))
        data = fetch_page(session_ref, page, page_size=page_size)

        if not data.get("success"):
            log.error("API вернул success=false: %s", data.get("error"))
            break

        items = data.get("items") or []
        pages = data.get("pages") or {}
        total = int(pages.get("total", 0) or 0)
        per_page = int(pages.get("per_page", page_size) or page_size)
        total_pages = (total + per_page - 1) // per_page if total else 0

        if first_iteration:
            log.info("всего страниц: %s (per_page=%s, total=%s)",
                     total_pages, per_page, total)
            first_iteration = False

        new_on_page = 0
        for it in items:
            row = item_to_row(it)
            key = row["url"] or row["title"]
            if not key or key in seen:
                continue
            seen.add(key)
            rows.append(row)
            new_on_page += 1

        save_csv(rows)
        save_xlsx(rows)
        save_progress(next_page=page + 1, total_pages=total_pages)
        log.info("страница %s/%s: новых %s, всего %s",
                 page, total_pages, new_on_page, len(rows))

        if not items or new_on_page == 0:
            log.info("новых нет — завершаю")
            break

        if total_pages and page >= total_pages:
            log.info("последняя страница достигнута")
            break

        page += 1
        time.sleep(2.0 + random.uniform(0, 1.5))

    save_csv(rows)
    save_xlsx(rows)
    return rows


def main() -> int:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
    )
    try:
        rows = scrape_all()
    except KeyboardInterrupt:
        log.warning("остановлено пользователем — прогресс сохранён")
        return 130
    log.info("готово: %s вакансий", len(rows))
    return 0


if __name__ == "__main__":
    sys.exit(main())